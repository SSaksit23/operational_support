"""Write the officer's chosen values back into a copy of the name list.

Only the name list is ours to edit; supplier PDFs (land confirmation, e-ticket)
are reported back to the supplier instead. Each changed cell is highlighted and
gets a comment saying what it was, what it is now, and who chose it.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

from .parsers.namelist import FLIGHT_RE, HEADER_KEYS, _norm_header

CHANGED = PatternFill("solid", fgColor="FFF2A8")
FLIGHT_VALUE_RE = re.compile(r"^(\S+) (\d{2}) (\w{3}) ([A-Z]{3})-([A-Z]{3}) (\d{2}:\d{2})-(\d{2}:\d{2})$")  # verify._fmt


@dataclass
class Fix:
    kind: str  # "flight" | "name"
    value: str  # flight: "PN6430 26 Oct BKK-CKG 17:55-22:15"; name: "SURNAME/GIVEN"
    by: str
    passport_no: str | None = None  # for kind == "name"


def apply_fixes(src: str | Path, fixes: list[Fix], out_path: str | Path) -> list[str]:
    """Returns one line per changed cell."""
    wb = openpyxl.load_workbook(io.BytesIO(Path(src).read_bytes()))
    ws = wb.worksheets[0]
    changes: list[str] = []
    header_row = next((c.row for row in ws.iter_rows() for c in row if str(c.value or "").strip() == "No."), None)
    cols = {}
    if header_row:
        for c in ws[header_row]:
            key = HEADER_KEYS.get(_norm_header(str(c.value or "")))
            if key and key not in cols:
                cols[key] = c.column

    def mark(cell, old, new, by):
        cell.value = new
        cell.fill = CHANGED
        cell.comment = Comment(f"แก้จาก: {old}\nเป็น: {new}\nเลือกโดย: {by}", "Ops agent")
        changes.append(f"{ws.title}!{cell.coordinate}: {old} → {new}")

    for fx in fixes:
        if fx.kind == "flight":
            m = FLIGHT_VALUE_RE.match(fx.value)
            if not m:
                continue
            fno, day, mon, o, d, dep, arr = m.groups()
            for row in ws.iter_rows():
                for c in row:
                    s = str(c.value or "")
                    fm = FLIGHT_RE.search(s)
                    # same flight, or the same leg (route) flown by another flight number
                    if not fm or not (fm.group(2) == fno or (fm.group(3), fm.group(4)) == (o, d)):
                        continue
                    dd, mm, yyyy = fm.group(1).split("/")
                    when = datetime.strptime(f"{day} {mon} {yyyy}", "%d %b %Y").strftime("%d/%m/%Y")
                    new_leg = f"{when} | {fno} | {o}-{d} | {dep}-{arr}"
                    if fm.group(0) != new_leg:
                        mark(c, s, s[:fm.start()] + new_leg + s[fm.end():], fx.by)
        elif fx.kind == "name" and fx.passport_no and header_row and "passport_no" in cols:
            surname, _, given = fx.value.partition("/")
            for r in range(header_row + 1, ws.max_row + 1):
                if str(ws.cell(row=r, column=cols["passport_no"]).value or "").strip().upper() == fx.passport_no.upper():
                    for key, new in (("surname", surname.strip()), ("given_name", given.strip())):
                        if key in cols and new:
                            cell = ws.cell(row=r, column=cols[key])
                            old = str(cell.value or "").strip()
                            if old.upper() != new.upper():
                                mark(cell, old, new, fx.by)
                    break
    wb.save(out_path)
    return changes
