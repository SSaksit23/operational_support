"""Fill the TL expense workbook from verified facts, with a company selector.

The selector works in two places:
  * in the app: the officer picks a company per document before generating;
  * inside Excel: a "เลือกบริษัท" sheet holds one dropdown per document, and
    the company headers are formulas that read the choice from a hidden
    company table. Switching the dropdown rewrites the header and tax block.
    Logos are pictures and cannot follow a formula, so a logo is placed only
    when the app generates the file for that company.

Company header slots (from the template):
  เบิกค่าทิป!B1:B4       tip payment voucher header
  เบิกเบี้ยเลี้ยง!B1:B4   TL allowance payment voucher header
  ใบปะหน้า!B25           tax invoice details for receipts (gift purchases)
"""

from __future__ import annotations

import hashlib
import io
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from .companies import Company
from .message import th_range
from .models import GroupBundle

SELECTOR_SHEET = "เลือกบริษัท"
TABLE_SHEET = "_companies"
MAX_LINES = 4
TITLES_TH = {"MR": "นาย", "MRS": "นาง", "MISS": "นางสาว", "MS": "นางสาว", "MSTR": "ด.ช.", "MASTER": "ด.ช."}

# document role -> (label on the selector sheet, header target)
ROLES = {
    "tip_voucher": ("ใบจ่ายเงินค่าทิปไกด์ (เบิกค่าทิป)", ("เบิกค่าทิป", "B1")),
    "allowance_voucher": ("ใบจ่ายเงินเบี้ยเลี้ยง TL (เบิกเบี้ยเลี้ยง)", ("เบิกเบี้ยเลี้ยง", "B1")),
    "tax_invoice": ("ข้อมูลออกใบกำกับภาษี (ใบปะหน้า)", ("ใบปะหน้า", "B25")),
}


HEADER_WIDTH = 58  # visible characters that fit across the voucher header at the template's font size


def _visible_len(text: str) -> int:
    """Thai combining marks take no width."""
    return sum(1 for ch in text if not ("\u0e31" == ch or "\u0e34" <= ch <= "\u0e3a" or "\u0e47" <= ch <= "\u0e4e"))


def fit_header(lines: list[str], max_lines: int = MAX_LINES, width: int = HEADER_WIDTH) -> list[str]:
    """Re-break a company block so each line fits the voucher header, without changing any words."""
    lines = list(lines)
    for _ in range(10):
        long_i = next((i for i, ln in enumerate(lines) if _visible_len(ln) > width), None)
        if long_i is None:
            break
        if len(lines) >= max_lines:
            pairs = [(i, _visible_len(lines[i]) + 1 + _visible_len(lines[i + 1])) for i in range(len(lines) - 1)
                     if long_i not in (i, i + 1)]
            pairs = [(i, n) for i, n in pairs if n <= width]
            if not pairs:
                break
            i = min(pairs, key=lambda x: x[1])[0]
            lines[i:i + 2] = [lines[i] + " " + lines[i + 1]]
            continue
        ln = lines[long_i]
        cut = max((m for m in range(len(ln)) if ln[m] == " " and _visible_len(ln[:m]) <= width), default=None)
        if cut is None:
            break
        lines[long_i:long_i + 1] = [ln[:cut].rstrip(), ln[cut + 1:].strip()]
    return lines


def _sheet(wb, part: str):
    ws = next((ws for ws in wb.worksheets if part in ws.title), None)
    if ws is None:
        raise TemplateError(f"The expense template has no sheet named like {part!r}")
    return ws


def fill_expense_workbook(template: str | Path, bundle: GroupBundle, companies: dict[str, Company],
                          selection: dict[str, str], out_path: str | Path,
                          gift_budget_per_pax: Optional[float] = 35, recalc: bool = True) -> Path:
    """selection: role -> company id, for every role in ROLES."""
    nl = bundle.name_list
    if not nl or not nl.start_date:
        raise ValueError("A name list with travel dates is required")
    missing = [r for r in ROLES if r not in selection]
    if missing:
        raise ValueError(f"Choose a company for: {', '.join(missing)}")

    wb = openpyxl.load_workbook(io.BytesIO(Path(template).read_bytes()))

    # Hidden company table: label | line 1..4 | tax id
    order = list(companies.values())
    tbl = wb.create_sheet(TABLE_SHEET)
    tbl.append(["label"] + [f"line {i}" for i in range(1, MAX_LINES + 1)] + ["tax_id", "id"])
    for c in order:
        lines = (fit_header(c.lines) + [""] * MAX_LINES)[:MAX_LINES]
        tbl.append([c.label] + lines + [c.tax_id, c.id])
    tbl.sheet_state = "hidden"
    last = len(order) + 1

    # Visible selector sheet with one dropdown per document
    sel = wb.create_sheet(SELECTOR_SHEET)
    sel["A1"] = "เลือกบริษัทที่ออกเอกสาร (เปลี่ยนได้จากรายการ แล้วหัวกระดาษจะเปลี่ยนตาม)"
    sel["A1"].font = Font(bold=True)
    dv = DataValidation(type="list", formula1=f"='{TABLE_SHEET}'!$A$2:$A${last}", allow_blank=False)
    sel.add_data_validation(dv)
    sel.column_dimensions["A"].width = 44
    sel.column_dimensions["B"].width = 28
    role_cell = {}
    for i, (role, (label, _)) in enumerate(ROLES.items(), start=3):
        sel.cell(row=i, column=1, value=label)
        cell = sel.cell(row=i, column=2, value=companies[selection[role]].label)
        dv.add(cell)
        role_cell[role] = f"'{SELECTOR_SHEET}'!$B${i}"

    def line_formula(role: str, n: int) -> str:
        col = "BCDE"[n]
        match = f"MATCH({role_cell[role]},'{TABLE_SHEET}'!$A$2:$A${last},0)"
        return f"T(INDEX('{TABLE_SHEET}'!${col}$2:${col}${last},{match}))"

    for role in ("tip_voucher", "allowance_voucher"):
        ws = _sheet(wb, ROLES[role][1][0])
        for n in range(MAX_LINES):
            ws.cell(row=1 + n, column=2, value="=" + line_formula(role, n))
    cover = _sheet(wb, "ใบปะหน้า")
    parts = ",".join(line_formula("tax_invoice", n) for n in range(MAX_LINES))
    cover["B25"] = f'="ข้อมูลสำหรับออกบิลกำกับภาษี"&CHAR(10)&_xlfn.TEXTJOIN(CHAR(10),TRUE,{parts})'

    # Group facts
    customers = [p for p in nl.passengers if not p.is_tour_leader]
    tl = next((p for p in nl.passengers if p.is_tour_leader), None)
    days = (nl.end_date - nl.start_date).days + 1
    tip = _sheet(wb, "เบิกค่าทิป")
    if tl:
        title = TITLES_TH.get((tl.title or "").upper().rstrip("."), "")
        old_payee = str(tip["B8"].value or "")
        tip["B8"] = f"{title} {tl.given_name_th or tl.given_name} {tl.surname_th or tl.surname}".strip()
        if (tl.given_name_th or tl.given_name) not in old_payee.replace("  ", " "):
            tip["B9"] = None  # the previous payee's home address; the TL fills in their own
    tip["B11"] = nl.tour_code
    tip["B12"] = f"{nl.program_name} ({nl.program_code}) : " if nl.program_code else nl.program_name
    tip["B13"] = th_range(nl.start_date, nl.end_date)
    allow_payee_addr = _label_row(_sheet(wb, "เบิกเบี้ยเลี้ยง"), "ที่อยู่ผู้รับเงิน")
    _sheet(wb, "เบิกเบี้ยเลี้ยง").cell(row=allow_payee_addr, column=2).value = "=T(เบิกค่าทิป!B9)"  # was showing 0
    _set_right_of(tip, "Day", days)
    _set_right_of(tip, "Pax", len(customers))

    _set_right_of(cover, "Pax", len(nl.passengers), keep_formula=True)
    if gift_budget_per_pax is not None:
        _set_right_of(cover, "Budget", gift_budget_per_pax)

    allow = _sheet(wb, "เบิกเบี้ยเลี้ยง")
    allow_row = _label_row(allow, "ค่าเบี้ยเลี้ยง")
    tl_tip_row = _label_row(allow, ("ค่าทิปTL", "ค่าทิป"))
    buy_row = _label_row(allow, "ค่าซื้อของ")
    allow.cell(row=allow_row, column=8, value=days)
    allow.cell(row=tl_tip_row, column=8, value=len(customers))
    allow.cell(row=buy_row, column=10).value = None  # receipts total from the last group; TL fills after the trip

    # Pictures: keep a logo only on the sheet whose selected company owns it; drop everything else
    # (the template carries a previous payee's ID card copy).
    for role in ("tip_voucher", "allowance_voucher"):
        ws = _sheet(wb, ROLES[role][1][0])
        own_logo = companies[selection[role]].logo_sha256
        keep = []
        for img in ws._images:
            data = img._data()
            img.ref = io.BytesIO(data)  # _data() consumes the stream; give save() a fresh one
            if hashlib.sha256(data).hexdigest() == own_logo:
                keep.append(img)
        ws._images = keep

    out_path = Path(out_path)
    wb.save(out_path)
    if recalc:
        _recalculate(out_path)
    return out_path


class TemplateError(ValueError):
    """The uploaded workbook does not have a label the filler needs."""


def _label_row(ws, label) -> int:
    from .parsers.expense_template import _find_label
    c = _find_label(ws, label)
    if c is None:
        raise TemplateError(f"Could not find {label if isinstance(label, str) else ' / '.join(label)!r} on sheet "
                            f"{ws.title.strip()!r} of the expense template")
    return c.row


def _set_right_of(ws, label: str, value, keep_formula: bool = False) -> None:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and not c.value.startswith("=") and c.value.strip() == label:
                target = ws.cell(row=c.row, column=c.column + 1)
                if keep_formula and isinstance(target.value, str) and target.value.startswith("="):
                    return  # newer templates link this cell to another sheet
                target.value = value
                return
    raise TemplateError(f"Could not find {label!r} on sheet {ws.title.strip()!r} of the expense template")


def _recalculate(path: Path) -> None:
    """Let LibreOffice compute the formulas so the file shows values even in viewers that never recalculate."""
    soffice = shutil.which("soffice")
    if not soffice:
        return
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                        "--outdir", tmp, str(path)], check=True, capture_output=True, timeout=120)
        produced = Path(tmp) / (path.stem + ".xlsx")
        if produced.exists():
            shutil.copyfile(produced, path)
