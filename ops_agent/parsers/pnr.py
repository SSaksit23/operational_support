"""Reservation record (PNR) reader: a GDS screen pasted into .docx or .txt.

Reads the three line types we rely on:
  " 1.ARUNLERTCHEEVIN/ANAN  2.ARUNLERTCHEEWIN/JANPHEN"   passenger names
  "25.PN6430  G  SA24OCT  BKKCKG  RR24  1935  2340  E  --T3"   air segments
  "131.TN/847-2120297577/P1"   ticket numbers per passenger
Segment dates carry no year; the weekday picks the year nearest to ref_date.
"""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path

import docx

from ..models import DocType, FlightLeg, PNRDoc, PNRPassenger, SourceRef
from ..normalize import hhmm, parse_dowddmon

NAME_RE = re.compile(r"(?:^|\s)(\d{1,3})\.([A-Z][A-Z ]*?/[A-Z][A-Z ]*?)(?=\s{2,}|\s*\d{1,3}\.[A-Z]|\s*$)")
SEG_RE = re.compile(
    r"^\s*(\d{1,3})\.\s*([A-Z0-9]{2})\s*(\d{1,4})\s+([A-Z])\s+([A-Z]{2}\d{1,2}[A-Z]{3})\s+([A-Z]{3})([A-Z]{3})\s+"
    r"([A-Z]{2}\d{0,3})\s+(\d{4})\s+(\d{4})")
TN_RE = re.compile(r"TN/(\d{3})-?(\d{10})/P(\d{1,3})")


def read_lines(path: Path, fmt: str | None = None) -> tuple[list[str], date | None]:
    if (fmt or path.suffix.lower().lstrip(".")) == "docx":
        d = docx.Document(str(path))
        lines = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                lines.append("  ".join(c.text for c in row.cells))
        created = d.core_properties.modified or d.core_properties.created
        return lines, created.date() if created else None
    return path.read_text(encoding="utf-8", errors="replace").splitlines(), None


def looks_like_pnr(lines: list[str]) -> bool:
    text = "\n".join(lines)
    return bool(TN_RE.search(text) or SEG_RE.search(text) and NAME_RE.search(text))


def parse_pnr(path: str | Path, display_name: str | None = None, ref_date: date | None = None,
              fmt: str | None = None) -> PNRDoc:
    path = Path(path)
    fname = display_name or path.name
    lines, doc_date = read_lines(path, fmt)
    ref = ref_date or doc_date or date.today()

    passengers: dict[int, PNRPassenger] = {}
    segments: list[FlightLeg] = []
    for i, line in enumerate(lines, start=1):
        if not line.strip():
            continue
        src = SourceRef(file=fname, doc_type=DocType.PNR, location=f"line {i}", raw=line.strip())
        seg = SEG_RE.match(line)
        if seg:
            segments.append(FlightLeg(
                flight_no=seg.group(2) + seg.group(3), date=parse_dowddmon(seg.group(5), ref),
                origin=seg.group(6), destination=seg.group(7), status=seg.group(8),
                dep_time=hhmm(seg.group(9)), arr_time=hhmm(seg.group(10)), source=src))
            continue
        tn = TN_RE.search(line)
        if tn:
            pno = int(tn.group(3))
            if pno in passengers:
                passengers[pno].ticket_no = tn.group(1) + tn.group(2)
            else:
                passengers[pno] = PNRPassenger(pnr_no=pno, name="", ticket_no=tn.group(1) + tn.group(2), source=src)
            continue
        for m in NAME_RE.finditer(line):
            pno = int(m.group(1))
            name = " ".join(m.group(2).split())
            if pno in passengers:
                passengers[pno].name = name
                passengers[pno].source = src
            else:
                passengers[pno] = PNRPassenger(pnr_no=pno, name=name, source=src)

    return PNRDoc(
        passengers=sorted(passengers.values(), key=lambda p: p.pnr_no),
        segments=segments,
        source=SourceRef(file=fname, doc_type=DocType.PNR, location="whole document"),
    )
