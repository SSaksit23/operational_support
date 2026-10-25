"""TL expense workbook: read the values a previous group may have left behind.

Cells are found by their labels (e.g. the number to the right of "Pax"),
not by fixed coordinates, so small layout edits to the template still work.
Phase 3 writes into the same slots.
"""

from __future__ import annotations

import hashlib
import io
from pathlib import Path

import openpyxl

from ..models import DocType, ExpenseTemplateDoc, Slot, SourceRef


# slot name -> (sheet name contains, label text, how to find the value)
#   "right": first value to the right of the label on the same row
#   ("col", n): value n columns to the right of the label
SLOTS = {
    "tip_days": ("เบิกค่าทิป", "Day", "right"),
    "tip_pax": ("เบิกค่าทิป", "Pax", "right"),
    "tour_code": ("เบิกค่าทิป", "รายละเอียด", "right"),
    "period": ("เบิกค่าทิป", "วันพรีเซ้น", "right"),
    "payee": ("เบิกค่าทิป", "ชื่อผู้รับเงิน", "right"),
    "cover_pax": ("ใบปะหน้า", "Pax", "right"),
    "cover_budget_per_pax": ("ใบปะหน้า", "Budget", "right"),
    "allowance_days": ("เบิกเบี้ยเลี้ยง", "ค่าเบี้ยเลี้ยง", ("col", 6)),
    "tl_tip_people": ("เบิกเบี้ยเลี้ยง", ("ค่าทิปTL", "ค่าทิป"), ("col", 6)),
}


def parse_expense_template(path: str | Path, display_name: str | None = None) -> ExpenseTemplateDoc:
    path = Path(path)
    fname = display_name or path.name
    # In memory, so embedded pictures stay readable; formulas kept, not cached results.
    wb = openpyxl.load_workbook(io.BytesIO(path.read_bytes()))
    out = ExpenseTemplateDoc(source=SourceRef(file=fname, doc_type=DocType.EXPENSE_TEMPLATE, location="workbook"))
    for slot, (sheet_part, label, how) in SLOTS.items():
        ws = next((ws for ws in wb.worksheets if sheet_part in ws.title), None)
        if ws is None:
            continue
        cell = _find_label(ws, label)
        if cell is None:
            continue
        target = None
        if how == "right":
            target = next((c for c in ws[cell.row] if c.column > cell.column and c.value not in (None, "")), None)
        else:
            target = ws.cell(row=cell.row, column=cell.column + how[1])
        if target is not None:
            out.slots[slot] = Slot(name=slot, value=target.value, source=SourceRef(
                file=fname, doc_type=DocType.EXPENSE_TEMPLATE, location=f"{ws.title.strip()}!{target.coordinate}",
                raw=None if target.value is None else str(target.value)))
    for ws in wb.worksheets:
        lines = [str(ws.cell(row=r, column=2).value).strip() for r in range(1, 5)
                 if ws.cell(row=r, column=2).value not in (None, "")]
        if lines and "บริษัท" in lines[0]:
            out.company_headers[ws.title.strip()] = lines
        for img in ws._images:
            cell = f"{_col(img.anchor._from.col)}{img.anchor._from.row + 1}"
            out.images.append(Slot(name="image", value=hashlib.sha256(img._data()).hexdigest(), source=SourceRef(
                file=fname, doc_type=DocType.EXPENSE_TEMPLATE, location=f"{ws.title.strip()}!{cell}")))
    return out


def _col(index: int) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(index + 1)


FORM_COLUMNS = 11  # labels live in A–K; columns further right hold rate notes ("-ลงร้าน ค่าเบี้ยเลี้ยง 1200/วัน")


FORM_ONLY = {"ค่าเบี้ยเลี้ยง", "ค่าทิปTL", "ค่าทิป"}  # also written in the rate notes to the right


def _find_label(ws, label):
    """First cell containing the label (or the first of several alternatives that exists)."""
    for one in (label,) if isinstance(label, str) else label:
        for row in ws.iter_rows(max_col=FORM_COLUMNS if one in FORM_ONLY else None):
            for c in row:
                if isinstance(c.value, str) and not c.value.startswith("=") and one in c.value:
                    return c
    return None
