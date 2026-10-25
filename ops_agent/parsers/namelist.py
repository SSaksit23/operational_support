"""Name list reader (.xlsx exported from the booking system).

The header block (tour code, leader, totals, flights) sits above a table
whose header row starts with "No.". Columns are found by header name, not
by fixed letters, so a reordered export still reads correctly.
"""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl

from ..models import DocType, FlightLeg, NameListDoc, Passenger, SourceRef
from ..normalize import hhmm, norm_id, norm_phone, to_date

COLUMN_KEYS = {
    "no.": "no", "room": "room", "title": "title", "name": "given_name", "surname": "surname",
    "ชื่อ": "given_name_th", "นามสกุล": "surname_th", "gender": "gender", "age": "age", "dob": "dob",
    "citizenid": "citizen_id", "passport": "passport_no", "doi": "passport_issue",
    "doe": "passport_expiry", "tel.": "tel", "agency": "agency", "remark อื่นๆ": "remark",
}

EXTRA_KEYS = {"sex": "gender", "id": "citizen_id", "tel": "tel", "remark": "remark", "passportno": "passport_no",
              "dateofbirth": "dob", "dateofissue": "passport_issue", "dateofexpiry": "passport_expiry", "expiry": "passport_expiry"}


def _norm_header(text: str) -> str:
    """'d.o.e.' / 'D.O.E' / 'doe' -> 'doe'; 'remark อื่นๆ' -> 'remarkอื่นๆ'."""
    return re.sub(r"[.\s:]", "", text.strip().lower())


HEADER_KEYS = {_norm_header(k): v for k, v in COLUMN_KEYS.items()} | EXTRA_KEYS

FLIGHT_RE = re.compile(
    r"(\d{2}/\d{2}/\d{4})\s*\|\s*([A-Z0-9]{2}\d{1,4})\s*\|\s*([A-Z]{3})-([A-Z]{3})\s*\|\s*(\d{1,2}:\d{2})-(\d{1,2}:\d{2})")


def _ref(file: str, cell: str, raw=None) -> SourceRef:
    return SourceRef(file=file, doc_type=DocType.NAME_LIST, location=cell,
                     raw=None if raw is None else str(raw))


def parse_namelist(path: str | Path, display_name: str | None = None) -> NameListDoc:
    path = Path(path)
    fname = display_name or path.name
    with open(path, "rb") as fh:  # file object: openpyxl rejects names without .xlsx
        wb = openpyxl.load_workbook(fh, data_only=True)
    ws = wb.worksheets[0]
    sheet = ws.title

    header: dict = {"flights": []}
    header_row = None
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if v is None:
                continue
            s = str(v).strip()
            loc = f"{sheet}!{c.coordinate}"
            if s == "No.":
                header_row = c.row
            elif s.startswith("Leader"):
                header["tour_leader"] = (s.split(":", 1)[1].strip(), loc)
            elif s.startswith("TOTAL") and "PAX" in s.upper():
                header["pax_total"] = (int(re.search(r"\d+", s).group()), loc)
                nxt = _row_text_after(ws, c)
                if nxt:
                    header["pax_breakdown"] = nxt
            elif s.startswith("TOTAL") and "ROOM" in s.upper():
                header["rooms_total"] = (int(re.search(r"\d+", s).group()), loc)
                nxt = _row_text_after(ws, c)
                if nxt:
                    header["room_mix"] = ({k: int(n) for n, k in re.findall(r"(\d+)\s*([A-Z]{3})", nxt[0])}, nxt[1])
            elif s.startswith("วันที่เดินทาง"):
                ds = re.findall(r"\d{2}/\d{2}/\d{4}", s)
                if len(ds) == 2:
                    header["dates"] = (to_date(ds[0]), to_date(ds[1]), loc)
            elif FLIGHT_RE.search(s):
                m = FLIGHT_RE.search(s)
                header["flights"].append(FlightLeg(
                    flight_no=m.group(2), date=to_date(m.group(1)), origin=m.group(3), destination=m.group(4),
                    dep_time=hhmm(m.group(5)), arr_time=hhmm(m.group(6)), source=_ref(fname, loc, s)))
            elif re.search(r"\(([A-Z0-9]+-[A-Z0-9]+)\)\s*:\s*(\S+)", s) and header_row is None:
                m = re.search(r"^(.*)\(([A-Z0-9]+-[A-Z0-9]+)\)\s*:\s*(\S+)", s)
                header["program"] = (m.group(1).strip(), m.group(2), m.group(3), loc)
            elif header_row is None and re.fullmatch(r"[A-Z0-9 ]{3,}", s) and "operator" not in header:
                header["operator"] = (s, loc)

    if header_row is None:
        raise ValueError(f"{fname}: no 'No.' header row found; not a name list?")

    cols = {}
    for c in ws[header_row]:
        if c.value is not None:
            key = HEADER_KEYS.get(_norm_header(str(c.value)))
            if key and key not in cols:
                cols[key] = c.column

    passengers: list[Passenger] = []
    current_room = None
    current_agency = None
    current_booking = None
    tl_name = (header.get("tour_leader") or ("", ""))[0].upper()
    for r in range(header_row + 1, ws.max_row + 1):
        def val(key):
            col = cols.get(key)
            return ws.cell(row=r, column=col).value if col else None

        given, surname = val("given_name"), val("surname")
        if not given or not surname:
            continue
        # Room and agency cells are filled only on the first person of a room / booking.
        if val("room"):
            current_room = str(val("room")).strip()
        agency_cell = val("agency")
        if agency_cell:
            lines = [x.strip() for x in str(agency_cell).splitlines() if x.strip()]
            current_agency = lines[0] if lines else None
            bk = re.search(r"BK\d+", str(agency_cell))
            current_booking = bk.group() if bk else None
        given, surname = str(given).strip().upper(), str(surname).strip().upper()
        is_tl = bool(tl_name) and given in tl_name and surname in tl_name
        remark = val("remark")
        remark = None if remark is None or str(remark).strip() in ("", "-") else str(remark).strip()
        age = val("age")
        passengers.append(Passenger(
            no=int(val("no")) if val("no") is not None else None,
            room=current_room, title=val("title"), given_name=given, surname=surname,
            given_name_th=val("given_name_th"), surname_th=val("surname_th"), gender=val("gender"),
            age=int(age) if age is not None else None, dob=to_date(val("dob")),
            citizen_id=norm_id(val("citizen_id")), passport_no=(str(val("passport_no")).strip() if val("passport_no") else None),
            passport_issue=to_date(val("passport_issue")), passport_expiry=to_date(val("passport_expiry")),
            tel=norm_phone(val("tel")), remark=remark, is_tour_leader=is_tl,
            # The TL travels for the company, not under the booking listed above them.
            agency=None if is_tl else current_agency, booking_no=None if is_tl else current_booking,
            source=_ref(fname, f"{sheet}!row {r}"),
        ))

    tour_code = header.get("program", (None, None, None))[2]
    dates = header.get("dates")
    return NameListDoc(
        tour_code=tour_code or "",
        program_name=header.get("program", (None,))[0],
        program_code=header.get("program", (None, None))[1],
        operator=header.get("operator", (None,))[0],
        tour_leader=header.get("tour_leader", (None,))[0],
        pax_total=header.get("pax_total", (None,))[0],
        pax_breakdown=header.get("pax_breakdown", (None,))[0],
        rooms_total=header.get("rooms_total", (None,))[0],
        room_mix=header.get("room_mix", ({},))[0],
        start_date=dates[0] if dates else None,
        end_date=dates[1] if dates else None,
        flights=header["flights"],
        passengers=passengers,
        source=_ref(fname, sheet),
    )


def _row_text_after(ws, cell):
    """First non-empty value to the right of a cell on the same row."""
    for c in ws[cell.row]:
        if c.column > cell.column and c.value not in (None, ""):
            return str(c.value).strip(), f"{ws.title}!{c.coordinate}"
    return None
