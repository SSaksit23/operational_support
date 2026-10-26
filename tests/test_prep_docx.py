"""A preparation letter sent as a Word file (made-up data)."""
from datetime import date

import docx
from docx.enum.text import WD_BREAK

from ops_agent.detect import detect
from ops_agent.parsers.prep_letter import page_texts, parse_prep_letter


def _letter(path):
    d = docx.Document()
    d.add_paragraph("ใบนัดหมายและเตรียมตัวเดินทาง")
    d.add_paragraph("คืนที่เข้าพัก ชื่อโรงแรม ที่อยู่")
    t = d.add_table(rows=0, cols=3)
    r = t.add_row().cells
    r[0].text, r[1].text, r[2].text = "24 ต.ค. 69", "TEST GRAND HOTEL\n(พัก 2 คืน)", "1 Test Road,\nTest City"
    d.add_paragraph("ป้ายต้อนรับ ณ สนามบินสุวรรณภูมิ")
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    d.add_paragraph("26 ต.ค. 69")
    d.add_paragraph("วันที่สอง เที่ยวเมือง")
    d.save(path)
    return path


def test_word_prep_letter(tmp_path):
    f = _letter(tmp_path / "ใบนัดหมาย TEST.docx")
    assert detect(f)[0].value == "prep_letter"
    assert len(page_texts(f)) == 2
    hotels = parse_prep_letter(f).hotels
    assert [(h.date, h.hotel_name, h.nights, h.address) for h in hotels] == [
        (date(2026, 10, 24), "TEST GRAND HOTEL", 2, "1 Test Road, Test City")]


def test_ocr_thread_limit_is_best_effort(monkeypatch):
    import onnxruntime as ort
    from ops_agent.parsers import passport
    monkeypatch.setattr(ort, "SessionOptions", ort.SessionOptions)
    passport._limit_threads(2)
    assert ort.SessionOptions().intra_op_num_threads == 2
